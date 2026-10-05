#!/usr/bin/env python3
"""
بيحوّل شيت المنتجات (SIMAT_English_PDP_Content_System.xlsx) لكتالوج الموقع:
src/lib/catalog.data.ts — ومنه بيتولّد supabase/seed.sql (npm run gen:seed).

    pip install openpyxl
    python3 scripts/import-catalog.py ../docs/data/SIMAT_English_PDP_Content_System.xlsx
    npm run gen:seed
"""
import json
import re
import sys
from pathlib import Path

import openpyxl

SRC = Path(sys.argv[1] if len(sys.argv) > 1 else '../docs/data/SIMAT_English_PDP_Content_System.xlsx')
OUT = Path(__file__).resolve().parent.parent / 'src' / 'lib' / 'catalog.data.ts'

# الشيت مفيهوش عمود Category (FOR HER / FOR HIM)، فاتحدد من العطر
# المستوحى منه. راجعه وعدّله هنا أو من لوحة التحكم.
GENDER = {
    **{i: 'women' for i in [1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 11, 12, 13, 14, 15, 35, 55, 56, 57]},
    **{i: 'men' for i in [20, 21, 23, 26, 27, 28, 29, 30, 33, 34, 36, 37, 38, 39, 40, 41, 43, 45, 46, 53]},
    **{i: 'unisex' for i in [16, 17, 18, 19, 22, 24, 25, 31, 32, 42, 44, 47, 48, 49, 50, 51, 52, 54, 58]},
}

# العطور اللي بتظهر في «الأكثر طلباً» والـ Mega Menu (بالترتيب).
FEATURED = [1, 19, 29, 34, 17, 55, 36, 8]


def slugify(name: str) -> str:
    name = re.sub(r'^SIMAT\s+', '', name.strip(), flags=re.I)
    return re.sub(r'[^a-z0-9]+', '-', name.lower()).strip('-')


def split_list(value) -> list[str]:
    return [p.strip() for p in str(value or '').split(',') if p.strip()]


def label_style(gender: str, family: str) -> str:
    if gender == 'men':
        return 'noir'
    if gender == 'unisex':
        return 'velvet'
    return 'bordeaux' if 'Amber' in family else 'linen'


wb = openpyxl.load_workbook(SRC, data_only=True)
ws = wb.worksheets[0]
header = [str(c.value).strip() if c.value else '' for c in ws[1]]
col = {name: i for i, name in enumerate(header) if name}

products = []
for row in ws.iter_rows(min_row=2, values_only=True):
    if not row[col['Product ID']]:
        continue
    pid = int(row[col['Product ID']])
    get = lambda key: str(row[col[key]] or '').strip()
    name = get('SIMAT Name')
    family = get('Fragrance Family')
    gender = GENDER.get(pid, 'unisex')
    featured_rank = FEATURED.index(pid) if pid in FEATURED else None
    products.append({
        'id': slugify(name),
        'slug': slugify(name),
        'sheetId': pid,
        'nameEn': name,
        'secondaryLine': get('Secondary Line'),
        'tagline': get('Tagline'),
        'shortDescription': get('Short Description'),
        'descriptionEn': get('Full Description'),
        'scentCharacter': get('Scent Character'),
        'familyEn': family,
        'accords': split_list(get('Main Accords')),
        'topNotesEn': split_list(get('Top Notes')),
        'heartNotesEn': split_list(get('Heart Notes')),
        'baseNotesEn': split_list(get('Base Notes')),
        'wearProfile': get('The Scent'),
        'occasion': get('Occasion / Character'),
        'related': [slugify(get(f'Related Product {n}')) for n in (1, 2, 3) if get(f'Related Product {n}')],
        'gender': gender,
        'labelStyle': label_style(gender, family),
        'isFeatured': featured_rank is not None,
        'soldCount': (100 - featured_rank * 5) if featured_rank is not None else max(1, 60 - pid),
    })

slugs = {p['slug'] for p in products}
for p in products:
    missing = [r for r in p['related'] if r not in slugs]
    if missing:
        print(f"⚠ {p['nameEn']}: related not found {missing}", file=sys.stderr)

OUT.write_text(
    '// مولّد تلقائياً من شيت المنتجات بـ scripts/import-catalog.py — متعدّلش بإيدك.\n'
    'export const catalogRows = '
    + json.dumps(products, ensure_ascii=False, indent=2)
    + ' as const;\n',
    encoding='utf-8',
)
print(f'✓ {len(products)} products → {OUT}')
